# spike_docx.py  (Phase 2 library spike — throwaway, outside the package)
#
# Goal: prove python-docx's native comment API and docx-revisions tracked
# changes can both live in the SAME .docx without corrupting each other.
#
# Workflow:
#   1. Build a .docx with python-docx, add one paragraph + a comment on it.
#   2. Add a tracked insertion via docx-revisions on the same document.
#   3. Save spike.docx.
#   4. Re-open with python-docx, print the comment text + author.
#   5. Convert spike.docx -> PDF via the package's headless wrapper.
#
# Run from the repo root:  uv run python scripts/spike_docx.py
# (PYTHONPATH must be stripped so the project venv's deps are used.)
from __future__ import annotations

import os
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
os.environ.pop("PYTHONPATH", None)

from docx import Document  # noqa: E402
from docx_revisions import RevisionParagraph  # noqa: E402
from libreoffice_mcp.config import load_settings  # noqa: E402
from libreoffice_mcp.utils import libreoffice as lo  # noqa: E402

OUT = REPO / "spike.docx"


def main() -> int:
    # ---- 1. Build + comment -------------------------------------------------
    doc = Document()
    para = doc.add_paragraph("Spike paragraph: base text.")
    run = para.runs[0]
    doc.add_comment(
        runs=para.runs,
        text="This is a spike comment.",
        author="Spike",
        initials="SP",
    )
    print("added comment on paragraph")

    # ---- 2. Tracked insertion via docx-revisions ---------------------------
    rp = RevisionParagraph.from_paragraph(para)
    tracked = rp.add_tracked_insertion(" TRACKED INSERTION", author="Spike")
    print(f"added tracked insertion text={tracked.text!r} author={tracked.author!r}")

    # ---- 3. Save ------------------------------------------------------------
    doc.save(str(OUT))
    print(f"saved {OUT.name}")

    # ---- 4. Re-open + read back the comment ---------------------------------
    reopened = Document(str(OUT))
    comments = list(reopened.comments)
    print(f"reopened: {len(comments)} comment(s)")
    if comments:
        c = comments[0]
        print(f"  comment text={c.text!r} author={c.author!r} initials={c.initials!r}")
    else:
        print("  !! NO COMMENTS READ BACK")
        return 1

    # ---- 5. Convert to PDF via headless wrapper -----------------------------
    settings = load_settings()
    pdf = lo.convert_file(OUT, "pdf", OUT.parent, settings, timeout=45.0)
    print(f"pdf conversion OK -> {pdf.name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())